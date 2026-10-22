# Copyright (c) 2026, X-SHA and contributors
# For license information, please see license.txt
#
# Import file "Fact_Detail_Sales" (detail penjualan per baris produk dari POS
# lama, 1 file per bulan per departemen, mis. 2026_01_FHS.xlsx) ke doctype
# Sales Invoice History. Semua 21 kolom file disimpan apa adanya; tambahan:
#   tanggal_transaksi = dari No Transaksi (CS/55/260101/0001 -> 2026-01-01),
#                       krn Waktu Bayar bisa jauh sesudahnya (piutang) / kosong
#   kode_outlet       = kode outlet ERP dari nama outlet di file
#   departemen        = FHS / FMCG dari nama file
#   sumber_file, baris_file = asal baris (nama dokumen = <file>-<baris>, jadi
#                       import ulang 1 file mengganti datanya, tdk dobel)
# Dijalankan di server (file terlalu besar utk GitHub), contoh:
#   bench --site <site> execute custom_reports.sales_invoice_history.import_folder \
#       --kwargs "{'path': '/home/frappe/Fact_Detail_Sales_2026'}"

import datetime
import os
import re

import frappe
from frappe.utils import now

DOCTYPE = "Sales Invoice History"
HEADER = ("No Transaksi", "Waktu Bayar", "Outlet", "Pelanggan", "Nama Produk", "SKU", "Kategori", "Quantity", "HPP Produk",
          "Harga Jual (Rp)", "Sub Total", "Nama Diskon Produk", "Diskon per Produk (Rp)", "Nama Diskon Transaksi",
          "Total Diskon (Rp)", "Komplimen (Rp)", "Kategori Transaksi", "Total Penjualan (Rp)", "Metode Pembayaran",
          "Status Pembayaran", "Jumlah Refund (Rp)")
FILE_FIELDS = ("no_transaksi", "waktu_bayar", "outlet", "pelanggan", "nama_produk", "sku", "kategori", "quantity", "hpp_produk",
               "harga_jual", "sub_total", "nama_diskon_produk", "diskon_per_produk", "nama_diskon_transaksi", "total_diskon",
               "komplimen", "kategori_transaksi", "total_penjualan", "metode_pembayaran", "status_pembayaran", "jumlah_refund")
NUMERIC = {"quantity", "hpp_produk", "harga_jual", "sub_total", "diskon_per_produk", "total_diskon", "komplimen",
           "total_penjualan", "jumlah_refund"}
OUTLET_CODE = {
    "X-Sha Cikiray": "XSC", "X-Sha Fashion SP": "XSP", "XSHA FURNITURE": "XJB", "XSHA PANYINGKIRAN": "XPY",
    "Xsha Ciawi": "XCW", "Xsha Pasar Singaparna": "PSR", "Tasmu Mart": "TASMU", "X-Sha Menara": "MNR",
    "X-Sha Supermart SP": "XSM", "X-Shamart Cisaruni": "XCR", "X-Shamart Mangunreja": "XMJ",
    "X-Shamart Puspahiang": "XPH", "X-Shamart Tanjungjaya": "XTJ", "X-Shamart Warung Peuteuy": "XWP",
}
_TRX_DATE = re.compile(r"^[A-Za-z]+/\d+/(\d{2})(\d{2})(\d{2})/\d+$")
_INSERT_FIELDS = ("name", "creation", "modified", "owner", "modified_by", "docstatus", "idx") + FILE_FIELDS + \
    ("tanggal_transaksi", "kode_outlet", "departemen", "sumber_file", "baris_file")


def _text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _num(v):
    if v in (None, ""):
        return 0
    return v if isinstance(v, (int, float)) else float(str(v).replace(",", ""))


def _waktu(v):
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, str) and v.strip():
        return datetime.datetime.strptime(v.strip(), "%d-%m-%Y %H:%M:%S")
    return None


def parse_rows(file_path):
    """Generator baris siap simpan dari 1 file: (baris_file, {field: nilai})."""
    from openpyxl import load_workbook

    fname = os.path.basename(file_path)
    departemen = "FMCG" if "FMCG" in fname.upper() else ("FHS" if "FHS" in fname.upper() else "")
    ws = load_workbook(file_path, read_only=True, data_only=True).worksheets[0]
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == 1:
            got = tuple(_text(x).strip() for x in row[:len(HEADER)])
            if got != HEADER:
                frappe.throw(f"{fname}: judul kolom tdk sesuai format Fact_Detail_Sales: {got}")
            continue
        if row[0] in (None, ""):
            continue
        d = {}
        for f, v in zip(FILE_FIELDS, row):
            if f == "waktu_bayar":
                d[f] = _waktu(v)
            elif f in NUMERIC:
                d[f] = _num(v)
            else:
                d[f] = _text(v)
        m = _TRX_DATE.match(d["no_transaksi"].strip())
        d["tanggal_transaksi"] = datetime.date(2000 + int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else \
            (d["waktu_bayar"].date() if d["waktu_bayar"] else None)
        d["kode_outlet"] = OUTLET_CODE.get(d["outlet"].strip(), "")
        d["departemen"] = departemen
        d["sumber_file"] = fname
        d["baris_file"] = i
        yield i, d


def import_file(file_path, chunk=5000):
    """Import 1 file (data lama dari file yg sama dihapus dulu). Return jumlah baris."""
    fname = os.path.basename(file_path)
    stem = os.path.splitext(fname)[0]
    frappe.db.delete(DOCTYPE, {"sumber_file": fname})
    frappe.db.commit()
    ts, user = now(), frappe.session.user or "Administrator"
    buf, n, unknown = [], 0, set()
    for i, d in parse_rows(file_path):
        if not d["kode_outlet"]:
            unknown.add(d["outlet"])
        buf.append((f"{stem}-{i:07d}", ts, ts, user, user, 0, 0) + tuple(d[f] for f in _INSERT_FIELDS[7:]))
        n += 1
        if len(buf) >= chunk:
            frappe.db.bulk_insert(DOCTYPE, _INSERT_FIELDS, buf)
            frappe.db.commit()
            buf = []
    if buf:
        frappe.db.bulk_insert(DOCTYPE, _INSERT_FIELDS, buf)
        frappe.db.commit()
    saved = frappe.db.count(DOCTYPE, {"sumber_file": fname})
    print(f"{fname}: {n} baris dibaca, {saved} tersimpan" + (f" | outlet tanpa kode: {sorted(unknown)}" if unknown else ""))
    if saved != n:
        frappe.throw(f"{fname}: jumlah tersimpan ({saved}) != dibaca ({n})")
    return n


def import_folder(path):
    """Import semua file .xlsx di folder (urut nama). Contoh lihat atas file."""
    files = sorted(f for f in os.listdir(path) if f.lower().endswith(".xlsx") and not f.startswith("~$"))
    if not files:
        frappe.throw(f"Tdk ada file .xlsx di {path}")
    total = 0
    for f in files:
        total += import_file(os.path.join(path, f))
    print(f"SELESAI: {len(files)} file, {total} baris | total di {DOCTYPE}: {frappe.db.count(DOCTYPE)}")
    return total
