# Copyright (c) 2026, X-SHA and contributors
# For license information, please see license.txt
#
# Download Excel laporan harian & mingguan dgn TAMPILAN SPT FILE ASLI admin (header
# bertingkat di-merge, warna, font, format angka), angka = persis hasil
# execute() report yg sama (tdk ada hitungan baru di sini). Export bawaan ERP
# tetap ada; ini tombol tambahan "Download Report".
#
# Weekly Margin -> meniru "WEEKLY MARGIN OCTOBER 2026" sheet "28-04 October"
# bagian WEEKLY MARGIN CLEAR (kolom N:Z di file, di sini mulai kolom A).
# Weekly Sales Achievement & Weekly Visitor TC AC -> meniru SALES REPORT HO
# sheet "Data weekly" (Mingguan = bagian kiri A:AI, Pembanding = kanan AK:BN).
# Weekly Sales Net -> "WEEKLY MARGIN CLEAR" Sheet4; Kontribusi -> "KONTRIBUSI
# SALES CATEGORY W1 VS W4" (1 sheet per departemen).
# Daily Sales vs Target -> SALES REPORT HO sheet ALL / DEPT. F&H.S / DEPT. FMCG;
# Daily Inventory -> file inventory harian (3 varian berdampingan).
# Laporan bulanan -> file admin masing2 (SALES REPORT HO, SALES CLEAR, SALES &
# MARGIN, VISITOR TC AC, INV TARGET, KONTRIBUSI MoM, SERVEL, SALES YTD); lihat bagian bawah.

import calendar
from io import BytesIO

import frappe
from frappe.utils import formatdate, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_ACC = '_-* #,##0_-;\\-* #,##0_-;_-* "-"_-;_-@_-'  # ribuan, 0 tampil "-" (spt file)
_PCT = "0.00%"
_PCT_RED = "0.00%;[Red]\\-0.00%"
_THIN = Side(style="thin")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _font(bold=False, size=10, color=None, name="Times New Roman"):
    return Font(name=name, bold=bold, size=size, color=color)


def _fill(rgb):
    return PatternFill("solid", fgColor=rgb) if rgb else PatternFill()


def _put(ws, row, col, value, bold=False, size=10, fill=None, fmt=None, align=None, color=None, border=True, rotate=0, wrap=False,
         font="Times New Roman"):
    c = ws.cell(row=row, column=col, value=value)
    c.font = _font(bold, size, color, font)
    if fill:
        c.fill = _fill(fill)
    if fmt:
        c.number_format = fmt
    if align or rotate or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center", text_rotation=rotate, wrap_text=wrap)
    if border:
        c.border = _BOX
    return c


def _send(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    frappe.response["filename"] = filename
    frappe.response["filecontent"] = buf.getvalue()
    frappe.response["type"] = "binary"


def _pct(v):
    return v / 100 if isinstance(v, (int, float)) else None


def build_weekly_margin(data, week_start, week_end):
    """Workbook Weekly Margin dari hasil execute() report Weekly Margin."""
    wb = Workbook()
    ws = wb.active
    ws.title = f"{formatdate(week_start, 'dd')}-{formatdate(week_end, 'dd MMMM')}"[:31]
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLM", (3.9, 4.7, 25.3, 15.1, 15.7, 13.4, 12.4, 23.9, 15.3, 14.3, 9.6, 9.1, 9.1)):
        ws.column_dimensions[col].width = w

    # judul (file: N1 "WEEKLY MARGIN CLEAR X-SHA GROUP", N2 "Periode:", Q2 "28-04 October")
    ws.merge_cells("A1:M1")
    _put(ws, 1, 1, "WEEKLY MARGIN CLEAR X-SHA GROUP", bold=True, size=12, align="center", border=False)
    ws.row_dimensions[1].height = 35.25
    ws.merge_cells("A2:C2")
    _put(ws, 2, 1, "Periode:", bold=True, size=14, align="right")
    ws.merge_cells("D2:K2")
    _put(ws, 2, 4, f"{formatdate(week_start, 'dd')}-{formatdate(week_end, 'dd MMMM')}", bold=True, size=14, align="left")
    ws.row_dimensions[2].height = 19.5

    # header bertingkat (file baris 3-5)
    blue, grey = "FF1E4E79", "FF7F7F7F"
    for rng in ("A3:A5", "B3:B5", "C3:C5", "D3:F4", "G3:M4"):
        ws.merge_cells(rng)
    for col in range(1, 14):
        for r in (3, 4, 5):
            _put(ws, r, col, None, bold=True, fill=blue if col <= 6 else grey, color="FFFFFFFF", align="center")
    _put(ws, 3, 2, "NO.", bold=True, fill=blue, color="FFFFFFFF", align="center")
    _put(ws, 3, 3, "OUTLET", bold=True, fill=blue, color="FFFFFFFF", align="center")
    _put(ws, 3, 4, "EXCLUDE PROMO", bold=True, fill=blue, color="FFFFFFFF", align="center", wrap=True)
    _put(ws, 3, 7, "INCLUDE PROMO", bold=True, fill=grey, color="FFFFFFFF", align="center", wrap=True)
    for col, label in zip(range(4, 14), ("DEAL MARGIN", "SALES ACHV", "MARGIN", "PROMO-X", "RAFAKSI DISCOUNT - X",
                                         "DEAL MARGIN", "SALES ACHV", "MARGIN", "TARGET", "ACHV")):
        _put(ws, 5, col, label, bold=True, fill=blue if col <= 6 else grey, color="FFFFFFFF", align="center")
    ws.row_dimensions[5].height = 18.75
    ws.freeze_panes = "D6"

    # kolom nilai: (field, format, fill kolom)
    cols = [("deal_margin_clear", _ACC, None), ("sales_retur", _ACC, None), ("margin_pct", _PCT, "FFB4C6E7"),
            ("promo", _ACC, None), ("rafaksi", _ACC, None), ("deal_margin_plus", _ACC, None), ("sales_achv", _ACC, None),
            ("margin_pct_incl", _PCT, "FFD6DCE4"), ("target", _PCT, None), ("achv", _PCT_RED, "FFC8C8C8")]
    total_fill = "FFD9E2F3"

    rows = [r for r in data if not r.get("is_info")]
    try:
        cut = next(i for i, r in enumerate(rows) if not r)  # baris kosong = sebelum tabel mitra
    except StopIteration:
        cut = len(rows)
    main, mitra = rows[:cut], [r for r in rows[cut + 1:] if r and not r.get("is_header")]

    r = 6
    section_start, section_label = None, None

    def close_section(end_row):
        if section_start is None or end_row < section_start:
            return
        if end_row > section_start:
            ws.merge_cells(start_row=section_start, start_column=1, end_row=end_row, end_column=1)
        _put(ws, section_start, 1, section_label, bold=True, align="center", rotate=90)
        for rr in range(section_start + 1, end_row + 1):
            ws.cell(rr, 1).border = _BOX

    for row in main:
        if row.get("is_header"):
            close_section(r - 1)
            section_start, section_label = r, ("" if row["outlet"] == "HO" else row["outlet"])
            continue
        label = row.get("outlet") or ""
        if label == "X-Sha Kemitraan":
            # file: Kemitraan s/d TOTAL punya blok kiri sendiri (N30:N33, tanpa label)
            close_section(r - 1)
            section_start, section_label = r, ""
        is_dept_total = label in ("Total", "TOTAL")
        fill = total_fill if is_dept_total else None
        bold = bool(row.get("bold"))
        _put(ws, r, 2, row.get("no"), bold=False, size=11, fill=fill, align="center")
        _put(ws, r, 3, label, bold=bold, size=11, fill=fill, align="left", color="FF000000")
        for i, (field, fmt, colfill) in enumerate(cols, start=4):
            v = row.get(field)
            if fmt in (_PCT, _PCT_RED):
                v = _pct(v)
            _put(ws, r, i, v, bold=bold or colfill is not None, fill=fill or colfill, fmt=fmt)
        r += 1
    close_section(r - 1)

    # rincian Kemitraan per mitra (file: tabel LABA / SALES / % di bawah TOTAL)
    if mitra:
        r += 1
        _put(ws, r, 4, "LABA", bold=True, align="center", border=False)
        _put(ws, r, 5, "SALES", bold=True, align="center", border=False)
        r += 1
        for row in mitra:
            is_total = row.get("bold")
            fill = total_fill if is_total else None
            _put(ws, r, 3, "" if is_total else row.get("outlet"), bold=bool(is_total), fill=fill, border=False)
            _put(ws, r, 4, row.get("deal_margin_clear"), bold=bool(is_total), fill=fill, fmt=_ACC, border=False)
            _put(ws, r, 5, row.get("sales_retur"), bold=bool(is_total), fill=fill, fmt=_ACC, border=False)
            _put(ws, r, 6, _pct(row.get("margin_pct")), bold=bool(is_total), fill=fill, fmt="0.0%", border=False)
            r += 1
    return wb


@frappe.whitelist()
def download_weekly_margin(tanggal=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.weekly_margin_target_vs_achievement.weekly_margin_target_vs_achievement import execute
    from custom_reports.sales_api import _monday_week_bounds

    as_of = getdate(tanggal)
    week_start, week_end = (getdate(d) for d in _monday_week_bounds(as_of))
    _cols, data = execute({"tanggal": str(as_of)})
    wb = build_weekly_margin(data, week_start, week_end)
    _send(wb, f"WEEKLY MARGIN {formatdate(week_start, 'dd')}-{formatdate(week_end, 'dd MMM yyyy')}.xlsx")


# ---------------------------------------------------------------------------
# Weekly Sales Achievement / Weekly Visitor TC AC (SALES REPORT HO "Data weekly")
_ROBOTO = "Roboto"
_DW_SECTION = "FF203864"   # label departemen vertikal (accent1 -50%)
_DW_GROWTH = "FFF8CBAD"    # header Growth (accent2 +60%)
_DW_MTD = "FF1F4E78"       # header MTD (accent5 -50%)
_DW_TOTAL = "FFD9D9D9"     # total departemen
_DW_HOLDING = "FFD0CECE"   # HOLDING (tampilan Mingguan)
_DW_HOLDING_CMP = "FFB4C6E7"  # HOLDING (tampilan Pembanding)
_DW_CUR_HEAD, _DW_CUR = "FFFFD966", "FFFFF2CC"  # kolom bulan berjalan di Pembanding
_DW_DEPT_TOTALS = ("FASHION & H S", "FMCG")


def _dw_rows(data):
    """[(section_label, [rows])] dari data report (baris header = awal section)."""
    sections = []
    for row in data:
        if row.get("is_info"):
            continue
        if row.get("is_header"):
            sections.append((row["outlet"], []))
        elif sections:
            sections[-1][1].append(row)
    return sections


def _dw_row_style(label):
    """(bold, jenis total) utk 1 baris Data weekly."""
    if label == "HOLDING":
        return True, "holding"
    if label in _DW_DEPT_TOTALS:
        return True, "total"
    return label.startswith("Total Outlet"), None


def _dw_mingguan(ws, data, month_label, title, widths):
    """Bagian kiri sheet Data weekly (kolom A:AI): TC/Sales/AC/Visitor per minggu, growth, MTD."""
    info = next((r["outlet"] for r in data if r.get("is_info")), "")
    head = dict(bold=True, size=10, align="center", font=_ROBOTO)
    # A label, B NO., C OUTLET, lalu W1 W2 Growth W3 Growth W4 Growth MTD (masing2 4 kolom)
    blocks = [("WEEK-1", "w1_", None), ("WEEK-2", "w2_", None), ("Growth", "g2_", _DW_GROWTH), ("WEEK-3", "w3_", None),
              ("Growth", "g3_", _DW_GROWTH), ("WEEK-4", "w4_", None), ("Growth", "g4_", _DW_GROWTH), ("MTD", "mtd_", _DW_MTD)]
    subs = (("tc", "TC"), ("sales", "Sales"), ("ac", "AC"), ("visitor", "VISITOR"))
    last_col = 3 + 4 * len(blocks)
    sd = [p.strip() for p in info.split("|") if "s/d tgl" in p]
    ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=last_col)
    _put(ws, 1, 3, f"{title} {month_label}" + (f" ({sd[0]})" if sd else ""), bold=True, size=11, align="center",
         border=False, font=_ROBOTO)
    _put(ws, 2, 2, "NO.", **head)
    _put(ws, 2, 3, "OUTLET", **head)
    _put(ws, 3, 2, None, **head)
    _put(ws, 3, 3, None, **head)
    ws.merge_cells("B2:B3")
    ws.merge_cells("C2:C3")
    for bi, (label, _pre, fill) in enumerate(blocks):
        c0 = 4 + 4 * bi
        white = "FFFFFFFF" if fill == _DW_MTD else None
        for k in range(4):
            _put(ws, 2, c0 + k, label if k == 0 else None, fill=fill, color=white, **head)
            _put(ws, 3, c0 + k, subs[k][1], fill=fill, color=white, **head)
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 3)
    widths.update({1: 3.3, 2: 3.7, 3: 22.9})
    for bi in range(len(blocks)):
        for k, w in enumerate((7.7, 12.0, 9.0, 7.7)):
            widths[4 + 4 * bi + k] = w
    r = 4
    for section, rows in _dw_rows(data):
        start = r
        for row in rows:
            label = row.get("outlet") or ""
            bold, kind = _dw_row_style(label)
            fill = {"total": _DW_TOTAL, "holding": _DW_HOLDING}.get(kind)
            _put(ws, r, 2, row.get("no"), size=9, fill=fill, align="center", font=_ROBOTO)
            _put(ws, r, 3, label, size=9, fill=fill, align="left", font=_ROBOTO, bold=bold)
            for bi, (_label, pre, _f) in enumerate(blocks):
                for k, (field, _sub) in enumerate(subs):
                    v, fmt = row.get(pre + field), _ACC
                    if pre.startswith("g"):
                        v, fmt = _pct(v), ("0.0%" if field == "tc" else _PCT)
                    _put(ws, r, 4 + 4 * bi + k, v, size=8, fill=fill, fmt=fmt, bold=bold, font=_ROBOTO)
            r += 1
        _put(ws, start, 1, section, bold=True, size=10, fill=_DW_SECTION, color="FFFFFFFF", align="center", rotate=90,
             font=_ROBOTO)
        if r - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=r - 1, end_column=1)
    ws.freeze_panes = "D4"


def _dw_pembanding(ws, data, month_label, m1_label, m2_label, o, widths):
    """Bagian kanan sheet Data weekly: per minggu vs tahun lalu & 2 bulan sebelumnya, digeser `o` kolom."""
    head = dict(bold=True, size=10, align="center", font=_ROBOTO)
    # Pembanding: per minggu 7 kolom = LAST YEAR (bln ini, thn lalu, growth) + LAST 2 MONTH (m-2, m-1, vs m-2, vs m-1)
    bulan, tahun = month_label.rsplit(" ", 1)
    subs = [("cur", month_label, _ACC), ("ly", f"{bulan} {int(tahun) - 1}", _ACC), ("ly_growth", "growth", _PCT),
            ("m2", m2_label, _ACC), ("m1", m1_label, _ACC), ("vs_m2", f"vs {m2_label[:3].lower()}", _PCT),
            ("vs_m1", f"vs {m1_label[:3].lower()}", _PCT)]
    for rr in (1, 2, 3):
        _put(ws, rr, o + 1, "NO." if rr == 1 else None, **head)
        _put(ws, rr, o + 2, "OUTLET" if rr == 1 else None, **head)
    ws.merge_cells(start_row=1, start_column=o + 1, end_row=3, end_column=o + 1)
    ws.merge_cells(start_row=1, start_column=o + 2, end_row=3, end_column=o + 2)
    for w in range(4):
        c0 = o + 3 + 7 * w
        for k in range(7):
            _put(ws, 1, c0 + k, f"WEEK-{w + 1}" if k == 0 else None, bold=True, size=11, align="center", font=_ROBOTO)
            _put(ws, 2, c0 + k, "LAST YEAR" if k == 0 else ("LAST 2 MONTH" if k == 3 else None), **head)
            _put(ws, 3, c0 + k, subs[k][1], fill=_DW_CUR_HEAD if k == 0 else None, **head)
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 6)
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 2)
        ws.merge_cells(start_row=2, start_column=c0 + 3, end_row=2, end_column=c0 + 6)
    widths.update({o + 1: 3.7, o + 2: 22.9})
    for w in range(4):
        for k, wd in enumerate((14.6, 13.0, 7.4, 13.6, 13.0, 7.1, 7.4)):
            widths[o + 3 + 7 * w + k] = wd
    r = 4
    for _section, rows in _dw_rows(data):
        for row in rows:
            label = row.get("outlet") or ""
            bold, kind = _dw_row_style(label)
            fill = {"total": _DW_TOTAL, "holding": _DW_HOLDING_CMP}.get(kind)
            _put(ws, r, o + 1, row.get("no"), size=9, fill=fill, align="center", font=_ROBOTO)
            _put(ws, r, o + 2, label, size=9, fill=fill, align="left", font=_ROBOTO, bold=bold)
            for w in range(4):
                for k, (field, _l, fmt) in enumerate(subs):
                    v = row.get(f"p{w + 1}_{field}")
                    if fmt == _PCT:
                        v = _pct(v)
                    _put(ws, r, o + 3 + 7 * w + k, v, size=8, fmt=fmt, bold=bold, font=_ROBOTO,
                         fill=fill or (_DW_CUR if k == 0 else None))
            r += 1


def build_weekly_sales(data, month_label, m1_label="", m2_label="", title="DATA WEEKLY", data_pembanding=None):
    """Workbook Data weekly: bagian Mingguan (kiri) + kalau ada data_pembanding, bagian Pembanding di
    kanannya mulai kolom AK - persis susunan sheet "Data weekly" di file asli (1 sheet)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data weekly"
    ws.sheet_view.showGridLines = False
    widths = {}
    _dw_mingguan(ws, data, month_label, title, widths)
    if data_pembanding is not None:
        _dw_pembanding(ws, data_pembanding, month_label, m1_label, m2_label, 36, widths)  # 36 -> NO. di kolom AK
        widths[36] = 2.0
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[3].height = 21
    return wb


def _weekly_sales_download(report_module, filters, title, filename_prefix, with_pembanding):
    from custom_reports.sales_api import _BULAN_ID
    mod = frappe.get_module(report_module)
    _cols, data = mod.execute(dict(filters, tampilan="Mingguan"))
    data_cmp = mod.execute(dict(filters, tampilan="Pembanding"))[1] if with_pembanding else None
    month = _BULAN_ID.index(filters.get("bulan")) + 1
    year = int(filters.get("tahun"))
    m1 = month - 1 if month > 1 else 12
    m2 = m1 - 1 if m1 > 1 else 12
    wb = build_weekly_sales(data, f"{_BULAN_ID[month - 1]} {year}", _BULAN_ID[m1 - 1], _BULAN_ID[m2 - 1], title, data_cmp)
    _send(wb, f"{filename_prefix} {_BULAN_ID[month - 1]} {year}.xlsx")


@frappe.whitelist()
def download_weekly_sales_achievement(filters=None):
    frappe.only_for("System Manager")
    filters = frappe.parse_json(filters) if isinstance(filters, str) else dict(filters or {})
    _weekly_sales_download("custom_reports.custom_reports.report.weekly_sales_achievement.weekly_sales_achievement",
                           filters, "DATA WEEKLY", "WEEKLY SALES ACHIEVEMENT", with_pembanding=True)


@frappe.whitelist()
def download_weekly_visitor(filters=None):
    frappe.only_for("System Manager")
    filters = frappe.parse_json(filters) if isinstance(filters, str) else dict(filters or {})
    _weekly_sales_download("custom_reports.custom_reports.report.weekly_visitor_tc_ac_achievement.weekly_visitor_tc_ac_achievement",
                           filters, "DATA WEEKLY VISITOR, TC, AC", "WEEKLY VISITOR TC AC", with_pembanding=False)


# ---------------------------------------------------------------------------
# Weekly Sales Net Achievement (rincian perhitungan Weekly Margin) -> meniru "WEEKLY MARGIN CLEAR 28-4 SEPT" Sheet4
_AN = "Arial Narrow"
_SN_HEAD = "FFA5A5A5"   # header abu (Penjualan s/d SALES BRUTO 2)
_SN_KEY = "FFFBE4D5"    # kolom hasil (file: SALES ACHV)
_SN_DEPT = "FFD6DCE5"   # Total Dept
_SN_GRAND = "FF002060"  # GRAND TOTAL


def build_weekly_sales_net(data):
    """Workbook rincian perhitungan Weekly Margin (WEEKLY MARGIN CLEAR Sheet4) dari execute() Weekly Sales Net Achievement."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet4"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    info = next((r["outlet"] for r in data if r.get("is_info")), "")
    periode = info.split("|")[0].replace("Minggu:", "").strip()
    _put(ws, 1, 1, "SALES & MARGIN ACHIEVEMENT X-SHA GROUP", bold=True, size=16, border=False, font=_AN)
    _put(ws, 2, 1, f"Periode : {periode}", bold=True, size=11, align="left", border=False, font=_AN)
    # (field, judul, fill header, format)  - urutan kolom spt file E:T
    cols = [("penjualan", "Penjualan", _SN_HEAD, _ACC), ("member", "MEMBER-X", _SN_HEAD, _ACC),
            ("potongan", "Potongan Penjualan", _SN_HEAD, _ACC), ("promo", "Promo - X", _SN_HEAD, _ACC),
            ("rafaksi", "RAFAKSI DISCOUNT - X", _SN_HEAD, _ACC), ("sales_bruto_2", "SALES BRUTO 2", _SN_HEAD, _ACC),
            ("sales_terlapor", "SALES TERLAPOR", _SN_HEAD, _ACC), ("hpp", "HPP ERP", None, _ACC),
            ("retur", "RETUR PENJUALAN", None, _ACC), ("retur_hpp", "RETUR HPP", None, _ACC),
            ("sales_retur", "SALES - RETUR", None, _ACC), ("sales_achv", "SALES ACHV", _SN_KEY, _ACC),
            ("hpp_retur", "HPP - RETUR", None, _ACC), ("deal_margin_clear", "deal margin clear", None, _ACC),
            ("margin_pct", "%", None, _PCT), ("deal_margin", "deal margin ++", None, _ACC)]
    ws.merge_cells("A3:C3")
    for c in (1, 2, 3):
        _put(ws, 3, c, "Department" if c == 1 else None, bold=True, size=11, fill=_SN_HEAD, align="center", color="FF000000", font=_AN)
    for i, (_f, label, fill, _fmt) in enumerate(cols, start=4):
        _put(ws, 3, i, label, bold=True, size=11, fill=fill, align="center", wrap=True, font=_AN if fill == _SN_HEAD else _CAL)
    ws.row_dimensions[3].height = 30
    widths = (6.0, 3.9, 23.4, 15.9, 13.0, 13.0, 13.0, 13.0, 15.0, 15.0, 15.0, 13.0, 9.0, 15.9, 15.0, 15.1, 15.0, 9.0, 15.0)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    r = 4
    section_start, section_label = None, None
    extra = []

    def close_section(end_row):
        if section_start is None or end_row < section_start:
            return
        _put(ws, section_start, 1, section_label, bold=True, size=12, align="center", rotate=90, font=_AN, color="FF000000")
        if end_row > section_start:
            ws.merge_cells(start_row=section_start, start_column=1, end_row=end_row, end_column=1)

    rows = [x for x in data if not x.get("is_info")]
    for idx, row in enumerate(rows):
        if not row:  # baris kosong -> sisanya catatan di bawah tabel (kmt)
            extra = [x for x in rows[idx + 1:] if x]
            break
        if row.get("is_header"):
            close_section(r - 1)
            section_start, section_label = r, row["outlet"]
            continue
        label = row.get("outlet") or ""
        bold = bool(row.get("bold"))
        fill, color = None, "FF000000"
        if label.startswith("Total Dept"):
            fill = _SN_DEPT
        elif label in ("HO (Offline)", "GRAND TOTAL"):
            if section_start is not None:
                close_section(r - 1)
                section_start = None
            if label == "GRAND TOTAL":
                fill, color = _SN_GRAND, "FFFFFFFF"
        if section_start is None:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            _put(ws, r, 1, label, bold=True, size=12, fill=fill, align="left", font=_AN, color=color)
            for c in (2, 3):
                ws.cell(r, c).border = _BOX
        else:
            _put(ws, r, 2, row.get("no"), bold=True, size=11, fill=fill, align="center", font=_AN, color=color)
            _put(ws, r, 3, label, bold=bold, size=12, fill=fill, align="left", font=_AN, color=color)
        for i, (field, _l, colfill, fmt) in enumerate(cols, start=4):
            v = row.get(field)
            if fmt == _PCT:
                v = _pct(v)
            _put(ws, r, i, v, bold=bold, size=10, fill=fill or (colfill if colfill == _SN_KEY else None), fmt=fmt,
                 font="Arial", color=color)
        r += 1
    close_section(r - 1)
    # kmt (Kemitraan) & total + kmt di bawah tabel, spt file baris 34-35
    for row in extra:
        _put(ws, r, 3, row.get("outlet"), bold=bool(row.get("bold")), size=11, border=False, font=_AN)
        for i, (field, _l, _f, fmt) in enumerate(cols, start=4):
            if field in ("sales_retur", "sales_achv", "deal_margin_clear", "deal_margin"):
                _put(ws, r, i, row.get(field), bold=bool(row.get("bold")), size=10, fmt=fmt, border=False, font="Arial")
        r += 1
    ws.freeze_panes = "D4"
    return wb


@frappe.whitelist()
def download_weekly_sales_net(tanggal=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.weekly_sales_achievement_net_retur_dan_potongan.weekly_sales_achievement_net_retur_dan_potongan import execute
    from custom_reports.sales_api import _monday_week_bounds

    as_of = getdate(tanggal)
    week_start, week_end = (getdate(d) for d in _monday_week_bounds(as_of))
    _cols, data = execute({"tanggal": str(as_of)})
    _send(build_weekly_sales_net(data),
          f"WEEKLY MARGIN CLEAR {formatdate(week_start, 'dd')}-{formatdate(week_end, 'dd MMM yyyy')}.xlsx")


# ---------------------------------------------------------------------------
# Kontribusi Sales per Category -> meniru "KONTRIBUSI SALES CATEGORY W1 VS W4"
# (1 sheet per departemen, gaya pivot Excel, kolom pemisah kuning)
_CAL = "Calibri"
_KC_SEP = "FFFFFF00"


def build_kontribusi(sheets, week_no, prev_no):
    """sheets: [(nama sheet, data execute())]."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, data in sheets:
        ws = wb.create_sheet(name)
        b = dict(bold=True, size=11, border=False, font=_CAL)
        ws.merge_cells("B1:D1")
        ws.merge_cells("F1:H1")
        _put(ws, 1, 2, f"WEEK {week_no}", align="center", **b)
        _put(ws, 1, 6, f"WEEK {prev_no}", align="center", **b)
        _put(ws, 1, 10, "GROWTH %", **b)
        for i, label in ((1, "Row Labels"), (2, "Sum of Total (Qty)"), (3, "Sum of Total (Rp)"), (4, "Kontribusi %"),
                         (6, "Sum of Total (Qty)"), (7, "Sum of Total (Rp)"), (8, "Kontribusi %"), (10, f"Kontribusi W{week_no} VS W{prev_no}")):
            _put(ws, 2, i, label, **b)
        info = next((x["kategori"] for x in data if x.get("is_info")), "")
        _put(ws, 1, 1, info.split("|")[1].strip() if "|" in info else "", size=9, border=False, font=_CAL)
        r = 3
        for row in data:
            if row.get("is_info"):
                continue
            bold = bool(row.get("bold"))
            _put(ws, r, 1, row.get("kategori"), bold=bold, size=11, border=False, font=_CAL)
            for i, field, fmt in ((2, "qty", _ACC), (3, "sales", _ACC), (4, "kontribusi", _PCT), (6, "qty_prev", _ACC),
                                  (7, "sales_prev", _ACC), (8, "kontribusi_prev", _PCT), (10, "growth", _PCT)):
                v = row.get(field)
                if fmt == _PCT:
                    v = _pct(v)
                    if row.get("bold") and field == "growth":
                        v = None
                _put(ws, r, i, v, bold=bold, size=11, fmt=fmt, border=False, font=_CAL)
            r += 1
        for rr in range(1, r):
            for c in (5, 9):
                ws.cell(rr, c).fill = _fill(_KC_SEP)
        for col, w in zip("ABCDEFGHIJ", (22.1, 17.3, 16.3, 12.1, 1.4, 18.7, 17.9, 12.1, 1.6, 19.7)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B3"
    return wb


@frappe.whitelist()
def download_kontribusi(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.weekly_kontribusi_sales_per_category.weekly_kontribusi_sales_per_category import execute

    filters = frappe.parse_json(filters) if isinstance(filters, str) else dict(filters or {})
    week_no = int(str(filters.get("minggu") or "W4").upper().lstrip("W"))
    prev_no = week_no - 1 if week_no > 1 else 4
    sheets = []
    for dept in ("F&H.S", "FMCG", "XSL"):
        _cols, data = execute(dict(filters, departemen=dept))
        sheets.append((dept, data))
    _send(build_kontribusi(sheets, week_no, prev_no),
          f"KONTRIBUSI SALES CATEGORY W{week_no} VS W{prev_no} {filters.get('bulan')} {filters.get('tahun')}.xlsx")


# ---------------------------------------------------------------------------
# Daily Sales vs Target (SALES REPORT HO sheet ALL / DEPT. F&H.S / DEPT. FMCG)
_OPEN = "Open Sans"
_SVT_HEAD, _SVT_HEAD2 = "FF8FAADC", "FFB4C6E7"          # ALL: header (accent1 +40% / +60%)
_SVT_DHEAD, _SVT_DHEAD2 = "FFA6A6A6", "FFD9D9D9"        # sheet dept: header abu2
_SVT_NAVY = "FF002060"                                   # font baris Total / judul seksi
_SVT_STORE_TOTAL = "FFD9E1F2"                            # FMCG "Total <toko>"
_SVT_EXC, _SVT_EXC_DARK = "FF2E75B6", "FF1F4E79"         # TOTAL Exclude .. (accent5 -25% / -50%)
_SVT_DEPT_EXC = "FF2F5597"                               # TOTAL DEPT .. EXC (accent1 -25%)
_SVT_GRAND, _SVT_GRAND_BAL = "FF203864", "FFADB9CA"      # GRAND TOTAL (+ kolom balance)
_SVT_KEYS = ("target", "daily_tc", "daily_sales", "acc_tc", "acc_sales", "achv", "vs_standard", "lm_tc", "lm_sales",
             "vs_lm_tc", "vs_lm_sales", "ly_tc", "ly_sales", "vs_ly_tc", "vs_ly_sales", "myly_sales", "vs_myly",
             "target_90", "balance_90", "avg_90", "target_100", "balance_100", "avg_100")
_SVT_PCT = {"achv", "vs_standard"}
_SVT_GROWTH = {"vs_lm_tc", "vs_lm_sales", "vs_ly_tc", "vs_ly_sales", "vs_myly"}


def _svt_row_style(label, is_all):
    """(fill, warna font, bold) baris spt file."""
    if label == "GRAND TOTAL":
        return _SVT_GRAND, "FFFFFFFF", True
    if label.startswith("TOTAL Exclude") or label == "FMCG Exclude MBG":
        return _SVT_EXC, "FFFFFFFF", True
    if label in ("TOTAL DEPT. F&H.S ALL", "TOTAL DEPT.FMCG ALL"):
        return _SVT_EXC_DARK, "FFFFFFFF", True
    if label in ("TOTAL DEPT. F&H.S EXC HCW", "TOTAL DEPT.FMCG EXC UIN"):
        return _SVT_DEPT_EXC, "FFFFFFFF", True
    if label.startswith("Total ") and not label.startswith("Total Outlet") and not is_all:
        return _SVT_STORE_TOTAL, None, True
    if label in ("Total", "3.Mini Soccer") or label.startswith("Total Outlet"):
        return None, _SVT_NAVY if is_all else None, True
    return None, None, False


def _svt_sheet(ws, data, as_of, dept):
    from openpyxl.formatting.rule import CellIsRule, FormulaRule

    is_all = dept == "ALL"
    lab = 2 if is_all else 3           # kolom label (ALL: B, dept: C dgn No di B)
    v0 = lab + 1                        # kolom angka pertama (Target)
    last = v0 + len(_SVT_KEYS) - 1
    days = calendar.monthrange(as_of.year, as_of.month)[1]
    lm = calendar.month_name[(as_of.month - 2) % 12 + 1]
    ly = f"{calendar.month_name[as_of.month]} {as_of.year - 1}"
    head, head2 = (_SVT_HEAD, _SVT_HEAD2) if is_all else (_SVT_DHEAD, _SVT_DHEAD2)
    ws.sheet_view.showGridLines = False

    title = {"ALL": "X-SHA GROUP", "DEPT. F&H.S": "Dept. Fashion & Home supplies", "DEPT. FMCG": "Dept. FMCG"}[dept]
    _put(ws, 1, 2, title, bold=True, size=16 if is_all else 14, border=False, font=_OPEN)
    _put(ws, 2, 2, f"Month : {as_of.day} {calendar.month_name[as_of.month]} {as_of.year}", bold=True, size=11,
         fill="FFD0CECE" if is_all else "FFBDD7EE", border=False, font=_OPEN)
    if lab > 2:
        ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=lab)
    _put(ws, 3, 2, "SALES VS TARGET", bold=True, size=11, border=False, font=_OPEN)
    _put(ws, 4, v0 + 7, "Last Month :", bold=True, size=9, border=False, font=_OPEN)
    _put(ws, 4, v0 + 11, "Last Year :", bold=True, size=9, border=False, font=_OPEN)
    _put(ws, 4, v0 + 17, "Balance of the day to go :", bold=True, size=9, border=False, font=_OPEN)
    _put(ws, 4, v0 + 19, days - as_of.day, bold=True, size=9, fill="FFFFFF00", align="center", font=_OPEN)

    # header 2 baris (file baris 6-7 di ALL, 5-6 di sheet dept)
    h1, h2 = 5, 6
    hs = 9 if is_all else 8
    groups = [("Target", 1, head), ("Daily sales" if is_all else "Daily", 2, head), ("Accumulated", 2, head),
              ("Achv vs Target", 1, head), ("VS Standard Achv", 1, head), (f"MTD {lm}", 2, head2),
              ("vs Last month", 2, head2), (f"MTD {ly}", 2, head2), ("vs Last year", 2, head2),
              (f"Monthly sales {ly}", 1, head2), (f"{as_of.year} VS {as_of.year - 1}", 1, head2),
              (0.9, 1, head), ("Balance To Achv 90%", 1, head), ("Average Daily sales in order to achieve 90%", 1, head),
              (1, 1, head), ("Balance To Achv 100%", 1, head), ("Average Daily Sales in order to Acv 100%", 1, head)]
    for rr in (h1, h2):
        for cc in range(2, v0):
            _put(ws, rr, cc, None, fill=head, font=_OPEN)
    _put(ws, h1, 2, "Department" if is_all else "OUTLET", bold=True, size=hs, fill=head, align="center", font=_OPEN)
    ws.merge_cells(start_row=h1, start_column=2, end_row=h2, end_column=lab)
    col = v0
    for text, span, fill in groups:
        for cc in range(col, col + span):
            for rr in (h1, h2):
                _put(ws, rr, cc, None, fill=fill, font=_OPEN)
        _put(ws, h1, col, text, bold=True, size=hs, fill=fill, align="center", wrap=True, font=_OPEN,
             fmt="0%" if isinstance(text, (int, float)) else None)
        if span == 2:
            ws.merge_cells(start_row=h1, start_column=col, end_row=h1, end_column=col + 1)
            sub = ("Tc", "Sales") if col == v0 + 1 else ("TC", "SALES")
            for i, s in enumerate(sub):
                _put(ws, h2, col + i, s, bold=True, size=hs, fill=fill, align="center", font=_OPEN)
        else:
            ws.merge_cells(start_row=h1, start_column=col, end_row=h2, end_column=col)
        col += span
    ws.row_dimensions[h1].height = 16.5
    ws.row_dimensions[h2].height = 25.5

    r = first = h2 + 1
    for row in data[1:]:  # data[0] = baris info
        label = row.get("outlet") or ""
        if row.get("is_header") or (label and "target" not in row):   # judul seksi
            for cc in range(2, last + 1):
                _put(ws, r, cc, None, font=_OPEN)
            _put(ws, r, lab, label, bold=True, size=11 if is_all else 10, color=_SVT_NAVY, font=_OPEN)
            r += 1
            continue
        fill, color, bold = _svt_row_style(label, is_all)
        if not is_all:
            _put(ws, r, 2, int(row["no"]) if row.get("no") else None, bold=bold, size=9, fill=fill, color=color,
                 align="center", font=_OPEN)
        _put(ws, r, lab, label or None, bold=bold, size=10, fill=fill, color=color, font=_OPEN)
        for i, k in enumerate(_SVT_KEYS):
            v = row.get(k)
            if k in _SVT_PCT or k in _SVT_GROWTH:
                v, fmt = _pct(v), (_PCT if k in _SVT_PCT else "0%")
            else:
                fmt = _ACC
            f, fc = fill, color
            if fill and color and i >= _SVT_KEYS.index("target_90"):  # spt GRAND TOTAL file: kolom balance abu2
                f, fc = _SVT_GRAND_BAL, None
            _put(ws, r, v0 + i, v, bold=bold, size=8, fill=f, color=fc, fmt=fmt,
                 align="center" if fmt != _ACC else None, font=_OPEN)
        r += 1
    end = r - 1

    # catatan bawah spt file: Standard Achieve / MTD/tgl / Total hari
    _put(ws, r, v0 + 4, "Standard Achieve :", bold=True, size=8, font=_OPEN)
    _put(ws, r, v0 + 5, None, font=_OPEN)
    ws.merge_cells(start_row=r, start_column=v0 + 4, end_row=r, end_column=v0 + 5)
    _put(ws, r, v0 + 6, as_of.day / days, bold=True, size=8, fmt=_PCT, align="center", font=_OPEN)
    _put(ws, r + 1, v0 + 4, "MTD/tgl", bold=True, size=8, border=False, font=_OPEN)
    _put(ws, r + 1, v0 + 5, as_of.day, bold=True, size=8, fill="FFFFFF00", align="center", font=_OPEN)
    _put(ws, r + 2, v0 + 4, "Total hari", bold=True, size=8, border=False, font=_OPEN)
    _put(ws, r + 2, v0 + 5, days, bold=True, size=8, align="center", font=_OPEN)

    # conditional formatting spt file (Achv / VS Standard berwarna, growth & balance minus merah)
    L = get_column_letter
    yellow = dict(fill=PatternFill("solid", bgColor="FFFFEB9C"), font=Font(color="FF9C5700"))
    green = dict(fill=PatternFill("solid", bgColor="FFC6EFCE"), font=Font(color="FF006100"))
    red = dict(font=Font(color="FFFF0000"))
    if end >= first:
        # hanya sel berangka (baris judul seksi kosong tdk ikut berwarna)
        for ci, rules in ((v0 + 5, ((">1", yellow), (">0.9", green))),
                          (v0 + 6, ((">0", yellow), (">-0.1", green), ("<0", red)))):
            cl = L(ci)
            for cond, style in rules:
                ws.conditional_formatting.add(f"{cl}{first}:{cl}{end}", FormulaRule(
                    formula=[f"AND(ISNUMBER({cl}{first}),{cl}{first}{cond})"], stopIfTrue=cond != "<0", **style))
        for k in ("vs_lm_tc", "vs_lm_sales", "vs_ly_tc", "vs_ly_sales", "vs_myly", "balance_90", "avg_90", "balance_100", "avg_100"):
            cl = L(v0 + _SVT_KEYS.index(k))
            ws.conditional_formatting.add(f"{cl}{first}:{cl}{end}", CellIsRule(operator="lessThan", formula=["0"], **red))

    widths = ([1.9, 34.1] if is_all else [2.4, 3.4, 27.1]) + [12.9, 6.9, 12, 6.9, 15.7, 8, 8.5, 9, 13, 6, 7, 9, 13, 6, 7,
                                                              15, 7, 15.9, 14.7, 14, 14.7, 14.3, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = ws.cell(h2 + 1, v0)


def build_daily_sales_vs_target(sheets, as_of):
    """sheets = [(departemen, data hasil execute() Daily Sales vs Target)]."""
    wb = Workbook()
    wb.remove(wb.active)
    for dept, data in sheets:
        _svt_sheet(wb.create_sheet(dept), data, as_of, dept)
    return wb


@frappe.whitelist()
def download_daily_sales_vs_target(tanggal=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.daily_sales_vs_target.daily_sales_vs_target import execute

    as_of = getdate(tanggal)
    sheets = [(d, execute({"tanggal": str(as_of), "departemen": d})[1]) for d in ("ALL", "DEPT. F&H.S", "DEPT. FMCG")]
    _send(build_daily_sales_vs_target(sheets, as_of), f"SALES REPORT HO {formatdate(as_of, 'dd MMM yyyy')}.xlsx")


# ---------------------------------------------------------------------------
# Daily Inventory (file inventory harian admin: 3 varian berdampingan A / J / S,
# tiap varian tabel INVENTORY FASHION / FMCG / HO)
_INV_HEAD = "FFB4C6E7"
_INV_SEP = "FFD0CECE"
_INV_COLS = ("Saldo Awal", "Pembelian", "Penjualan", "Transfer Masuk", "Transfer Keluar", "Lainnya", "Saldo Akhir")
_INV_KEYS = ("saldo_awal", "pembelian", "penjualan", "transfer_masuk", "transfer_keluar", "lainnya", "saldo_akhir")


def build_daily_inventory(variants, sheet_title):
    """variants = [(label varian, data hasil execute() Daily Inventory)]."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.sheet_view.showGridLines = False
    last_row = 1
    for vi, (label, data) in enumerate(variants):
        c0 = 1 + vi * 9
        _put(ws, 1, c0, label, size=11, border=False, font=_CAL)
        r = 2
        for row in data:
            name = row.get("outlet")
            if not name:          # pemisah antar tabel
                r += 1
                continue
            if row.get("is_header"):
                for cc in range(c0, c0 + 8):
                    _put(ws, r, cc, None, fill=_INV_HEAD, font=_CAL)
                _put(ws, r, c0, name, bold=True, size=10, fill=_INV_HEAD, align="center", font=_CAL)
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 7)
                _put(ws, r + 1, c0, None, fill=_INV_HEAD, font=_CAL)
                for i, h in enumerate(_INV_COLS, 1):
                    _put(ws, r + 1, c0 + i, h, bold=True, size=10, fill=_INV_HEAD, align="center", font=_CAL)
                r += 2
                continue
            total = name == "TOTAL INVENTORY"
            fill = _INV_HEAD if total else None
            _put(ws, r, c0, name, bold=True, size=10, fill=fill, font=_CAL)
            for i, k in enumerate(_INV_KEYS, 1):
                _put(ws, r, c0 + i, row.get(k), bold=total, size=10 if total else 11, fill=fill, fmt=_ACC, font=_CAL)
            r += 1
        last_row = max(last_row, r - 1)
    for rr in range(1, last_row + 1):
        for cc in (9, 18):
            ws.cell(rr, cc).fill = _fill(_INV_SEP)
    for i in range(3):
        ws.column_dimensions[get_column_letter(1 + i * 9)].width = 25.7
        for j in range(1, 8):
            ws.column_dimensions[get_column_letter(1 + i * 9 + j)].width = 15
    for cc in ("I", "R"):
        ws.column_dimensions[cc].width = 2.4
    ws.freeze_panes = "B2"
    return wb


@frappe.whitelist()
def download_daily_inventory(from_date=None, to_date=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.daily_inventory.daily_inventory import execute

    from_date, to_date = getdate(from_date), getdate(to_date or from_date)
    variants = [(v, execute({"from_date": str(from_date), "to_date": str(to_date), "variant": v})[1])
                for v in ("ALL", "EXC TOPUP", "EXC TOPUP&SEM")]
    if from_date == to_date:
        period, title = formatdate(to_date, "dd MMM yyyy"), formatdate(to_date, "dd MMM").upper()
    else:
        period = f"{formatdate(from_date, 'dd MMM')} - {formatdate(to_date, 'dd MMM yyyy')}"
        title = f"{formatdate(from_date, 'dd MMM')}-{formatdate(to_date, 'dd MMM')}".upper()
    _send(build_daily_inventory(variants, title), f"INVENTORY {period}.xlsx")


# ===========================================================================
# LAPORAN BULANAN
# ===========================================================================
_MO_SEP = "FFC4C4C4"
_MO_RP = '_-"Rp"* #,##0_-;\\-"Rp"* #,##0_-;_-"Rp"* "-"_-;_-@_-'


def _filters(filters):
    return frappe.parse_json(filters) if isinstance(filters, str) else dict(filters or {})


def _month_of(filters):
    """(month, year, m1, y1) dari filter bulan/tahun, default bulan berjalan (spt report)."""
    from frappe.utils import today
    from custom_reports.sales_api import _BULAN_ID

    now = getdate(today())
    month = _BULAN_ID.index(filters.get("bulan")) + 1 if filters.get("bulan") in _BULAN_ID else now.month
    year = int(filters.get("tahun") or now.year)
    m1, y1 = (month - 1, year) if month > 1 else (12, year - 1)
    return month, year, m1, y1


def _mmm_yy(month, year):
    return f"{calendar.month_abbr[month]}-{str(year)[2:]}"


def _bulan(month, year, upper=True):
    from custom_reports.sales_api import _BULAN_ID
    s = f"{_BULAN_ID[month - 1]} {year}"
    return s.upper() if upper else s


def _red_if_negative(ws, ref):
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="FFFF0000")))


def _mo_table(ws, data, r, cols, lab_font, lab_size, val_font, val_size, total_fill, grand_fill=_SN_GRAND,
              sec_fill=None, sec_color="FF000000", total_labels=None, no_font=None):
    """Tabel departemen gaya file bulanan: kolom A label departemen vertikal
    (merge), B No, C outlet, nilai mulai kolom D. cols = [(field, format, fill
    kolom)] atau None (kolom pemisah abu2). Baris "Total Dept.."/total_labels
    ber-fill total_fill, GRAND TOTAL navy. Return baris terakhir + 1."""
    section = {"start": None, "label": None}

    def close(end_row):
        if section["start"] is None or end_row < section["start"]:
            return
        _put(ws, section["start"], 1, section["label"], bold=True, size=lab_size, fill=sec_fill, color=sec_color,
             align="center", rotate=90, font=lab_font)
        if end_row > section["start"]:
            ws.merge_cells(start_row=section["start"], start_column=1, end_row=end_row, end_column=1)
        section["start"] = None

    for row in data:
        if row.get("is_info"):
            continue
        if not row:
            close(r - 1)
            r += 1
            continue
        if row.get("is_header"):
            close(r - 1)
            section["start"], section["label"] = r, row["outlet"]
            continue
        label = row.get("outlet") or ""
        bold = bool(row.get("bold"))
        fill, color = None, "FF000000"
        if label == "GRAND TOTAL":
            fill, color = grand_fill, "FFFFFFFF"
        elif label.startswith("Total Dept") or label.startswith("DEPT") or label in (total_labels or ()):
            fill = total_fill
        if label in ("HO (Offline)", "GRAND TOTAL") or section["start"] is None:
            close(r - 1)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            _put(ws, r, 1, label, bold=True, size=lab_size, fill=fill, align="left", font=lab_font, color=color)
            for c in (2, 3):
                ws.cell(r, c).border = _BOX
        else:
            _put(ws, r, 2, row.get("no"), bold=True, size=lab_size - 1, fill=fill, align="center", font=no_font or lab_font,
                 color=color)
            _put(ws, r, 3, label, bold=bold, size=lab_size, fill=fill, align="left", font=lab_font, color=color)
        for i, spec in enumerate(cols, start=4):
            if spec is None:
                _put(ws, r, i, None, fill=_MO_SEP, font=val_font)
                continue
            field, fmt, colfill = spec
            v = row.get(field)
            if fmt in (_PCT, "0%", "0.0%"):
                v = _pct(v)
            _put(ws, r, i, v, bold=bold, size=val_size, fill=fill or colfill, fmt=fmt, font=val_font, color=color)
        r += 1
    close(r - 1)
    return r


def _mo_header(ws, r, cols, labels, fill, font, size, color="FF000000", height=30, sep_fill=_MO_SEP):
    for i, (spec, label) in enumerate(zip(cols, labels), start=4):
        _put(ws, r, i, label, bold=True, size=size, fill=sep_fill if spec is None else fill, align="center", wrap=True,
             font=font, color=color)
    ws.row_dimensions[r].height = height


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# --- 1. Monthly Sales Achievement vs Target (= Daily Sales vs Target akhir bulan)
@frappe.whitelist()
def download_monthly_sales_achievement(filters=None):
    frappe.only_for("System Manager")
    from frappe.utils import today
    from custom_reports.custom_reports.report.monthly_sales_achievement_vs_target.monthly_sales_achievement_vs_target import execute

    filters = _filters(filters)
    month, year, _m1, _y1 = _month_of(filters)
    now = getdate(today())
    as_of = now if (year, month) == (now.year, now.month) else \
        getdate(f"{year}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}")
    sheets = [(d, execute(dict(filters, departemen=d))[1]) for d in ("ALL", "DEPT. F&H.S", "DEPT. FMCG")]
    _send(build_daily_sales_vs_target(sheets, as_of), f"SALES REPORT HO {_bulan(month, year)}.xlsx")


# --- 2. Monthly Sales Achv vs Target Net Retur & Potongan -> "SALES CLEAR ACHV VS TARGET MoM YoY"
_SC_HEAD, _SC_HEAD2 = "FFA5A5A5", "FF969696"


def _rincian_sheet(ws, data, month, year, title, with_margin):
    """Sheet4 file SALES CLEAR / SALES & MARGIN: rincian Penjualan s/d SALES ACHV (+ HPP & margin)."""
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    _put(ws, 1, 1, title, bold=True, size=16, border=False, font=_AN)
    cols = [("target", _ACC, None), ("penjualan", _ACC, None), ("member", _ACC, None), ("potongan", _ACC, None),
            ("promo", _ACC, None), ("rafaksi", _ACC, None), ("sales_bruto_2", _ACC, None), ("sales_terlapor", _ACC, None)]
    labels = [f"{calendar.month_abbr[month].upper()} Target", "Penjualan", "MEMBER-X", "Potongan Penjualan", "Promo - X",
              "RAFAKSI DISCOUNT - X", "SALES BRUTO 2", "SALES TERLAPOR"]
    if with_margin:
        cols += [("hpp", _ACC, None), ("retur", _ACC, None), ("retur_hpp", _ACC, None), ("sales_retur", _ACC, None),
                 ("sales_achv", _ACC, _SN_KEY), ("hpp_retur", _ACC, None), ("deal_margin_clear", _ACC, None),
                 ("margin_pct", _PCT, None), ("deal_margin", _ACC, None)]
        labels += ["HPP ERP", "RETUR PENJUALAN", "RETUR HPP", "SALES - RETUR", "SALES ACHV", "HPP - RETUR",
                   "deal margin clear", "%", "deal margin ++"]
    else:
        cols += [("retur", _ACC, None), ("sales_retur", _ACC, None), ("sales_achv", _ACC, _SN_KEY)]
        labels += ["RETUR PENJUALAN", "SALES - RETUR", "SALES ACHV"]
    last = 3 + len(cols)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last)
    _put(ws, 3, 1, _mmm_yy(month, year), bold=True, size=11, fill="FFFFFF00", align="center", font=_AN)
    for c in range(2, last + 1):
        _put(ws, 3, c, None, fill="FFFFFF00", font=_AN)
    ws.merge_cells("A4:C4")
    for c in (1, 2, 3):
        _put(ws, 4, c, "Department" if c == 1 else None, bold=True, size=11, fill=_SC_HEAD, align="center", font=_AN)
    for i, ((_f, _fmt, colfill), label) in enumerate(zip(cols, labels), start=4):
        head = _SC_HEAD if i <= 11 else (colfill or None)
        _put(ws, 4, i, label, bold=True, size=11, fill=head, align="center", wrap=True, font=_AN if i <= 11 else _CAL)
    ws.row_dimensions[4].height = 30
    _mo_table(ws, data, 5, cols, _AN, 12, "Arial", 10, _SN_DEPT)
    _widths(ws, (6.0, 3.9, 25.4) + (15.9,) * len(cols))
    ws.freeze_panes = "D5"


def build_monthly_sales_net(ringkasan, rincian, month, year, m1, y1):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:C2")
    for r in (1, 2):
        for c in (1, 2, 3):
            _put(ws, r, c, "Department" if (r, c) == (1, 1) else None, bold=True, size=11, fill=_SC_HEAD, align="center", font=_AN)
    # (field, fmt, fill) - kolom D..M; tahun lalu cuma sales (histori tanpa target)
    cols = [("target", _ACC, None), ("sales_achv", _ACC, None), ("achv", "0%", None),
            ("target_lm", _ACC, None), ("sales_achv_lm", _ACC, None), ("achv_lm", "0%", None),
            ("sales_ly", _ACC, None), None, ("growth_lm", _PCT, None), ("growth_ly", _PCT, None)]
    groups = [(4, 6, _mmm_yy(month, year)), (7, 9, _mmm_yy(m1, y1)), (10, 10, _mmm_yy(month, year - 1)), (12, 13, "GROWTH")]
    for c0, c1, text in groups:
        for c in range(c0, c1 + 1):
            _put(ws, 1, c, None, fill=_SC_HEAD2, font=_CAL)
        _put(ws, 1, c0, text, bold=True, size=11, fill=_SC_HEAD2, align="center", font=_CAL)
        if c1 > c0:
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c1)
    _put(ws, 1, 11, None, fill=_MO_SEP, font=_CAL)
    _mo_header(ws, 2, cols, ["TARGET", "SALES ACHV", "SALES VS TARGET", "TARGET", "SALES ACHV", "SALES VS TARGET",
                             "SALES ACHV", None, "VS LAST MONTH", "VS LAST YEAR"], _SC_HEAD2, _CAL, 11, height=18)
    end = _mo_table(ws, ringkasan, 3, cols, _AN, 12, _CAL, 11, _SN_DEPT)
    for cl in ("L", "M"):
        _red_if_negative(ws, f"{cl}3:{cl}{end}")
    _widths(ws, (4.6, 3.9, 25.0, 15.3, 14.3, 16.3, 15.3, 14.3, 16.3, 15.3, 2.7, 15.4, 12.9))
    ws.freeze_panes = "D3"
    _put(ws, end + 1, 3, f"* {_mmm_yy(month, year - 1)}: sales histori (kotor), tanpa target", size=9, border=False, font=_CAL)
    _rincian_sheet(wb.create_sheet("Sheet4"), rincian, month, year, "SALES ACHIEVEMENT X-SHA GROUP", with_margin=False)
    return wb


@frappe.whitelist()
def download_monthly_sales_net(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.monthly_sales_achievement_vs_target_net_retur_dan_potongan.monthly_sales_achievement_vs_target_net_retur_dan_potongan import execute

    filters = _filters(filters)
    month, year, m1, y1 = _month_of(filters)
    ring = execute(dict(filters, tampilan="Ringkasan"))[1]
    rinc = execute(dict(filters, tampilan="Rincian"))[1]
    _send(build_monthly_sales_net(ring, rinc, month, year, m1, y1),
          f"SALES CLEAR ACHV VS TARGET MoM YoY {_bulan(month, year)}.xlsx")


# --- 5. Monthly Margin -> "SALES & MARGIN ACHV" Sheet1 + Sheet4
def build_monthly_margin(ringkasan, rincian, month, year, end_day):
    from openpyxl.formatting.rule import CellIsRule, FormulaRule

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.sheet_view.showGridLines = False
    tnr = "Times New Roman"
    ws.merge_cells("B1:N1")
    _put(ws, 1, 2, "MONTHLY SALES & MARGIN ACHV X-SHA GROUP", bold=True, size=14, align="center", border=False, font=tnr)
    ws.merge_cells("A2:C2")
    _put(ws, 2, 1, "Periode:", bold=True, size=14, align="right", border=False, font=tnr)
    from custom_reports.sales_api import _BULAN_ID
    _put(ws, 2, 4, f"01-{end_day:02d} {_BULAN_ID[month - 1]} {year}", bold=True, size=14, align="left", border=False, font=tnr)
    cols = [("target", _ACC, None), ("sales_achv", _ACC, None), ("achv", _PCT, None), None,
            ("deal_margin", _ACC, None), ("target_margin_pct", _PCT, None), ("margin_achv_pct", _PCT, None),
            ("achv_margin", _PCT, None), None, ("target_margin_rp", _ACC, None), ("margin_achv_rp", _ACC, None)]
    for c in (1, 2, 3):
        _put(ws, 3, c, {2: "NO.", 3: "OUTLET"}.get(c), bold=True, size=11, fill=_SN_GRAND, color="FFFFFFFF", font=_CAL)
    _mo_header(ws, 3, cols, ["TARGET SALES", "SALES ACHV", "ACHV VS TARGET", None, "DEAL MARGIN", "TARGET MARGIN %",
                             "MARGIN ACHV %", "ACHV VS TARGET", None, "TARGET MARGIN RP", "MARGIN ACHV RP"],
               _SN_GRAND, _CAL, 11, color="FFFFFFFF", height=18, sep_fill="FFAEABAB")
    end = _mo_table(ws, ringkasan, 4, [c if c else None for c in cols], _CAL, 11, _CAL, 11, _SN_DEPT)
    for r in range(4, end):
        for c in (7, 12):
            ws.cell(r, c).fill = _fill("FFAEABAB")
    green, red = Font(color="FF00B050"), Font(color="FFC00000")
    for cl in ("F", "K"):
        rng = f"{cl}4:{cl}{end - 1}"
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({cl}4),{cl}4>=1)"], font=green))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f"AND(ISNUMBER({cl}4),{cl}4<1)"], font=red))
    ws.conditional_formatting.add(f"N4:N{end - 1}", CellIsRule(operator="lessThan", formula=["0"], font=red))
    _widths(ws, (3.7, 4.4, 25.3, 15.3, 15.0, 17.7, 1.7, 15.3, 18.3, 16.3, 17.7, 1.1, 20.4, 18.6))
    ws.freeze_panes = "D4"
    _rincian_sheet(wb.create_sheet("Sheet4"), rincian, month, year, "SALES & MARGIN ACHIEVEMENT X-SHA GROUP", with_margin=True)
    return wb


@frappe.whitelist()
def download_monthly_margin(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.monthly_margin_target_vs_achievement.monthly_margin_target_vs_achievement import execute

    filters = _filters(filters)
    month, year, _m1, _y1 = _month_of(filters)
    ring = execute(dict(filters, tampilan="Ringkasan"))[1]
    rinc = execute(dict(filters, tampilan="Rincian"))[1]
    info = ring[0].get("outlet", "") if ring else ""
    try:
        end_day = int(info.split("tgl 1-")[1].split()[0])
    except (IndexError, ValueError):
        end_day = calendar.monthrange(year, month)[1]
    _send(build_monthly_margin(ring, rinc, month, year, end_day), f"SALES & MARGIN ACHV {_bulan(month, year)}.xlsx")


# --- 3. Monthly Visitor TC AC -> "VISITOR, TC, AC ACHV VS TARGET, MoM YoY"
_VIS_SEC = "FF223962"
_VIS_ACHV = "FFF8CBAD"


def build_monthly_visitor(data, month, year, m1, y1):
    wb = Workbook()
    ws = wb.active
    ws.title = "EXC MBG,KMT,XSL"
    ws.sheet_view.showGridLines = False
    rb = _ROBOTO
    cols = []
    for pre in ("", "lm_"):
        cols += [(f"{pre}t_tc", _ACC, None), (f"{pre}t_ac", _ACC, None), (f"{pre}t_visitor", _ACC, None),
                 (f"{pre}tc", _ACC, None), (f"{pre}ac", _ACC, None), (f"{pre}visitor", _ACC, None),
                 (f"{pre}achv_tc", "0%", None), (f"{pre}achv_ac", "0%", None), (f"{pre}achv_visitor", "0%", None)]
    cols += [("ly_tc", _ACC, None), ("ly_ac", _ACC, None), None,
             ("g_tc_lm", _PCT, None), ("g_ac_lm", _PCT, None), ("g_tc_ly", _PCT, None), ("g_ac_ly", _PCT, None)]
    # baris 1: periode, baris 2: TARGET/ACTUAL/ACHV, baris 3: TC/AC/VISITOR
    for c in (1, 2, 3):
        for r in (1, 2, 3):
            _put(ws, r, c, {(1, 2): "NO.", (1, 3): "OUTLET"}.get((r, c)), bold=True, size=10, align="center", font=rb)
    ws.merge_cells("B1:B3")
    ws.merge_cells("C1:C3")
    blocks = [(4, _bulan(month, year)), (13, _bulan(m1, y1))]
    for c0, text in blocks:
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 8)
        _put(ws, 1, c0, text, bold=True, size=9, align="center", font=rb)
        for j, (sub, fill) in enumerate((("TARGET", None), ("ACTUAL", None), ("ACHV VS TARGET", _VIS_ACHV))):
            cc = c0 + j * 3
            ws.merge_cells(start_row=2, start_column=cc, end_row=2, end_column=cc + 2)
            _put(ws, 2, cc, sub, bold=True, size=10, fill=fill, align="center", font=rb)
            for k, t in enumerate(("TC", "AC", "VISITOR")):
                _put(ws, 3, cc + k, t, bold=True, size=10, fill=fill, align="center", font=rb)
            for k in (1, 2):
                _put(ws, 2, cc + k, None, fill=fill, font=rb)
        for k in range(1, 9):
            _put(ws, 1, c0 + k, None, font=rb)
    ly0 = 22
    ws.merge_cells(start_row=1, start_column=ly0, end_row=2, end_column=ly0 + 1)
    _put(ws, 1, ly0, f"{_bulan(month, year - 1)} (histori)", bold=True, size=9, align="center", wrap=True, font=rb)
    for r, c in ((1, ly0 + 1), (2, ly0), (2, ly0 + 1)):
        _put(ws, r, c, None, font=rb)
    _put(ws, 3, ly0, "TC", bold=True, size=10, align="center", font=rb)
    _put(ws, 3, ly0 + 1, "AC", bold=True, size=10, align="center", font=rb)
    for r in (1, 2, 3):
        _put(ws, r, ly0 + 2, None, fill="FF808080", font=rb)
    g0 = ly0 + 3
    ws.merge_cells(start_row=1, start_column=g0, end_row=1, end_column=g0 + 3)
    _put(ws, 1, g0, "GROWTH", bold=True, size=10, align="center", font=rb)
    for (c, sub, sub2) in ((g0, "LAST MONTH (achv)", "TC"), (g0 + 1, None, "AC"), (g0 + 2, "LAST YEAR", "TC"), (g0 + 3, None, "AC")):
        _put(ws, 2, c, sub, bold=True, size=9, fill="FFFFD966", align="center", font=rb)
        _put(ws, 3, c, sub2, bold=True, size=10, fill="FFFFD966", align="center", font=rb)
    ws.merge_cells(start_row=2, start_column=g0, end_row=2, end_column=g0 + 1)
    ws.merge_cells(start_row=2, start_column=g0 + 2, end_row=2, end_column=g0 + 3)
    for k in range(1, 4):
        _put(ws, 1, g0 + k, None, font=rb)
    end = _mo_table(ws, data, 4, cols, rb, 9, rb, 8, "FFD9D9D9", sec_fill=_VIS_SEC, sec_color="FFFFFFFF",
                    total_labels=("FASHION & H S", "FMCG"))
    for r in range(4, end):
        ws.cell(r, ly0 + 2).fill = _fill("FF808080")
    for c in range(g0, g0 + 4):
        _red_if_negative(ws, f"{get_column_letter(c)}4:{get_column_letter(c)}{end - 1}")
    _widths(ws, (3.3, 3.7, 22.9) + (9.5,) * 18 + (10.5, 10.5, 1.5, 11, 11, 11, 11))
    ws.freeze_panes = "D4"
    return wb


@frappe.whitelist()
def download_monthly_visitor(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.monthly_visitor_tc_ac_achievement_vs_target.monthly_visitor_tc_ac_achievement_vs_target import execute

    filters = _filters(filters)
    month, year, m1, y1 = _month_of(filters)
    data = execute(filters)[1]
    _send(build_monthly_visitor(data, month, year, m1, y1),
          f"VISITOR, TC, AC ACHV VS TARGET, MoM YoY {_bulan(month, year)}.xlsx")


# --- 4. Monthly Inventory Target vs Inventory Akhir -> "INV TARGET VS INV AKHIR" sheet EXC KMT, MBG, XSL
def build_monthly_inventory(data):
    from openpyxl.formatting.rule import FormulaRule

    wb = Workbook()
    ws = wb.active
    ws.title = "EXC KMT, MBG, XSL"
    ws.sheet_view.showGridLines = False
    _put(ws, 1, 1, "EXC TOPUP", size=11, border=False, font=_CAL)
    heads = ("OUTLET", "TARGET INV", "INV PLUS", "INV MINS", "TOTAL INV", "VS TARGET", "PERSENTASE")
    keys = ("target", "plus", "minus", "total", "vs_target", "pct")
    for i, h in enumerate(heads, start=1):
        _put(ws, 2, i, h, bold=True, size=9, fill=_MO_SEP, align="center", font=_OPEN)
    r = 3
    for row in data:
        if row.get("is_info"):
            continue
        if not row:
            r += 1
            continue
        label = row["outlet"]
        bold = bool(row.get("bold"))
        fill = "FFD9E1F2" if label.startswith("DEPT") or label == "HO" else None
        _put(ws, r, 1, label, bold=bold, size=9, fill=fill, font=_OPEN)
        for i, k in enumerate(keys, start=2):
            v = row.get(k)
            _put(ws, r, i, _pct(v) if k == "pct" else v, bold=bold, size=9, fill=fill, fmt="0%" if k == "pct" else _ACC,
                 align="center" if k == "pct" else None, font=_OPEN)
        r += 1
    rng = f"G3:G{r - 1}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=["AND(ISNUMBER(G3),G3<0.9)"], font=Font(color="FFFFC000")))
    ws.conditional_formatting.add(rng, FormulaRule(formula=["AND(ISNUMBER(G3),G3>1)"], font=Font(color="FFC00000")))
    info = next((x["outlet"] for x in data if x.get("is_info")), "")
    note = [p.strip() for p in info.split("|") if "blm diinput" in p]
    if note:
        _put(ws, r + 1, 1, f"* {note[0]}", size=9, border=False, font=_OPEN, color="FFC00000")
    _widths(ws, (25.4, 18.9, 16, 15.6, 16, 17.6, 12.1))
    ws.freeze_panes = "B3"
    return wb


@frappe.whitelist()
def download_monthly_inventory(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.monthly_inventory_target_vs_inventory_akhir.monthly_inventory_target_vs_inventory_akhir import execute

    filters = _filters(filters)
    month, year, _m1, _y1 = _month_of(filters)
    _send(build_monthly_inventory(execute(filters)[1]), f"INV TARGET VS INV AKHIR {_bulan(month, year)}.xlsx")


# --- 6. Monthly Kontribusi Sales per Category -> "KONTRIBUSI SALES CATEGORY AGUSTUS MoM YoY"
_KM_SEP = "FFADADAD"


def build_monthly_kontribusi(sheets, month, year, m1, y1):
    wb = Workbook()
    wb.remove(wb.active)
    for dept, data in sheets:
        ws = wb.create_sheet(dept)
        ws.sheet_view.showGridLines = False
        for c0, text in ((2, _bulan(month, year)), (6, _bulan(m1, y1))):
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + 2)
            _put(ws, 1, c0, text, bold=True, size=11, align="center", border=False, font=_CAL)
        _put(ws, 1, 10, "KOMPARASI", bold=True, size=11, align="center", border=False, font=_CAL)
        heads = ["CATEGORY", "Sum of Qty Terjual", "Sum of Sales (Rp)", "Sum of Kontribusi (%)", None,
                 "Sum of Qty Terjual", "Sum of Sales (Rp)", "Sum of Kontribusi (%)", None, "VS LAST MONTH"]
        for i, h in enumerate(heads, start=1):
            _put(ws, 2, i, h, bold=True, size=11, border=False, font=_CAL)
        keys = ["kategori", "qty", "sales", "kontribusi", None, "qty_lm", "sales_lm", "kontribusi_lm", None, "growth_lm"]
        r = 3
        for row in data:
            if row.get("is_info") or not row:
                continue
            bold = bool(row.get("bold"))
            for i, k in enumerate(keys, start=1):
                if k is None:
                    continue
                v = row.get(k)
                fmt = _PCT if k in ("kontribusi", "kontribusi_lm", "growth_lm") else (_ACC if k != "kategori" else None)
                if fmt == _PCT:
                    v = _pct(v)
                if row.get("kategori") == "Grand Total" and k == "growth_lm":
                    v = None
                _put(ws, r, i, v, bold=bold, size=11, fmt=fmt, border=False, font=_CAL)
            r += 1
        for rr in range(1, r):
            for c in (5, 9):
                ws.cell(rr, c).fill = _fill(_KM_SEP)
        _red_if_negative(ws, f"J3:J{r - 1}")
        _widths(ws, (22.1, 17.6, 16.4, 20.4, 2.4, 17.6, 16.4, 20.4, 2.6, 15.4))
        ws.freeze_panes = "B3"
        _put(ws, r + 1, 1, f"* {_bulan(month, year - 1)} (tahun lalu) tdk tersedia: data penjualan per item di ERP "
                            f"baru ada sejak akhir April 2026", size=9, border=False, font=_CAL)
    return wb


@frappe.whitelist()
def download_monthly_kontribusi(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.monthly_kontribusi_sales_per_category.monthly_kontribusi_sales_per_category import execute

    filters = _filters(filters)
    month, year, m1, y1 = _month_of(filters)
    sheets = [(d, execute(dict(filters, departemen=d))[1]) for d in ("F&H.S", "FMCG", "XSL")]
    _send(build_monthly_kontribusi(sheets, month, year, m1, y1),
          f"KONTRIBUSI SALES CATEGORY {_bulan(month, year)} MoM.xlsx")


# --- 7. Monthly Service Level per Category -> "SERVEL <BULAN> FMCG" (ITEM, SERVEL KATEGORI, supplier)
def _servel_pivot(ws, data):
    ws.sheet_view.showGridLines = False
    for i, h in enumerate(("Row Labels", "Sum of TOTAL NILAI PO (RP)", "Sum of TOTAL NILAI RCV (RP)", "SERVEL"), start=1):
        _put(ws, 1, i, h, bold=True, size=11, fill="FFDDEBF7", border=False, font=_CAL, align="right" if i == 4 else None)
    r = 2
    for row in data:
        if row.get("is_info"):
            continue
        grand = row.get("label") == "Grand Total"
        fill = "FFFFFF00" if grand else None
        _put(ws, r, 1, row.get("label"), bold=grand, size=11, fill=fill, border=False, font=_CAL)
        _put(ws, r, 2, row.get("nilai_po"), bold=grand, size=11, fill=fill, fmt=_MO_RP, border=False, font=_CAL)
        _put(ws, r, 3, row.get("nilai_rcv"), bold=grand, size=11, fill=fill, fmt=_MO_RP, border=False, font=_CAL)
        _put(ws, r, 4, _pct(row.get("servel")), bold=grand, size=11, fill=fill, fmt=_PCT, border=False, font=_CAL)
        r += 1
    _widths(ws, (40, 29.7, 30.7, 9.5))
    ws.freeze_panes = "A2"


def build_monthly_service_level(kategori, supplier, items):
    wb = Workbook()
    ws = wb.active
    ws.title = "ITEM"
    ws.sheet_view.showGridLines = False
    rows = [x for x in items if not x.get("is_info") and x.get("label") != "Grand Total"]
    grand = next((x for x in items if x.get("label") == "Grand Total"), {})
    # baris 1 = total (file: SUM), baris 2 = judul kolom
    for c, k, fmt in ((4, "po_qty", _ACC), (5, "rcv_qty", _ACC), (7, "nilai_po", _MO_RP), (8, "nilai_rcv", _MO_RP)):
        _put(ws, 1, c, grand.get(k), bold=True, size=11, fmt=fmt, border=False, font=_CAL)
    _put(ws, 1, 10, _pct(grand.get("servel")), bold=True, size=11, fmt="0%", border=False, font=_CAL)
    heads = ("ID", "ITEM NAME", "KATEGORI", "TOTAL PO (QTY)", "TOTAL RCV (QTY)", "LAST PURCHASE RATE", "TOTAL NILAI PO (RP)",
             "TOTAL NILAI RCV (RP)", "supplier", "SERVEL")
    for i, h in enumerate(heads, start=1):
        _put(ws, 2, i, h, bold=True, size=11, fill="FFDDEBF7", border=False, font=_CAL)
    keys = (("label", "@"), ("item_name", None), ("kategori", None), ("po_qty", _ACC), ("rcv_qty", _ACC),
            ("last_purchase_rate", _MO_RP), ("nilai_po", _MO_RP), ("nilai_rcv", _MO_RP), ("supplier_name", None), ("servel", "0%"))
    for r, row in enumerate(rows, start=3):
        for i, (k, fmt) in enumerate(keys, start=1):
            v = row.get(k)
            if k == "servel":
                v = _pct(v)
            c = ws.cell(row=r, column=i, value=v)
            if fmt:
                c.number_format = fmt
    _widths(ws, (15.1, 43.3, 25.1, 15.0, 16.0, 16.7, 19.1, 20.1, 34.0, 10))
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:J{max(len(rows) + 2, 3)}"
    _servel_pivot(wb.create_sheet("SERVEL KATEGORI"), kategori)
    _servel_pivot(wb.create_sheet("SERVEL SUPPLIER"), supplier)
    return wb


@frappe.whitelist()
def download_monthly_service_level(filters=None):
    frappe.only_for("System Manager")
    from custom_reports.custom_reports.report.monthly_service_level_per_category.monthly_service_level_per_category import execute

    filters = _filters(filters)
    month, year, _m1, _y1 = _month_of(filters)
    views = {t: execute(dict(filters, tampilan=t))[1] for t in ("Kategori", "Supplier", "Item")}
    dept = filters.get("departemen") or "FMCG"
    _send(build_monthly_service_level(views["Kategori"], views["Supplier"], views["Item"]),
          f"SERVEL {_bulan(month, year)} {dept}.xlsx")


# --- Sales YTD -> "SALES X-SHA GROUP_YTD_ 2026" (1 sheet per outlet / departemen / TOTAL)
_YTD_FONT = "Century Gothic"
_YTD_HEAD = "FFE7E6E6"


def _ytd_sheet(ws, title, sheet):
    from openpyxl.formatting.rule import FormulaRule
    from custom_reports.sales_ytd import has_kontribusi

    ws.sheet_view.showGridLines = False
    _put(ws, 1, 2, title, bold=True, size=10, border=False, font=_YTD_FONT)
    years = sheet["years"]
    _put(ws, 3, 2, "MONTH", bold=True, size=10, fill=_YTD_HEAD, align="left", font=_YTD_FONT)
    _put(ws, 4, 2, None, fill=_YTD_HEAD, font=_YTD_FONT)
    ws.merge_cells("B3:B4")
    col, layout = 3, []
    for i, y in enumerate(years):
        subs = [("target", "TARGET", _ACC), ("sales", "SALES", _ACC), ("ach", "ACH", "0.0%")]
        if i:
            subs.append(("growth", "GROWTH", _PCT))
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(subs) - 1)
        for j, (key, label, fmt) in enumerate(subs):
            _put(ws, 3, col + j, y if j == 0 else None, bold=True, size=11, fill=_YTD_HEAD, align="center", font=_YTD_FONT)
            _put(ws, 4, col + j, label, bold=True, size=10, fill=_YTD_HEAD, align="right" if key != "growth" else "center",
                 font=_YTD_FONT)
            layout.append((y, key, fmt, col + j))
        col += len(subs)
        if has_kontribusi(years, y):
            _put(ws, 3, col, "kontribusi", bold=True, size=10, fill=_YTD_HEAD, align="center", font=_YTD_FONT)
            _put(ws, 4, col, None, fill=_YTD_HEAD, font=_YTD_FONT)
            ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
            layout.append((y, "kontribusi", _PCT, col))
            col += 1

    def write(r, item, bold, label_size, label_align="left", val_size=8):
        _put(ws, r, 2, item["label"], bold=True, size=label_size, align=label_align, font=_YTD_FONT)
        for y, key, fmt, c in layout:
            v = item["per_year"][y][key]
            if key in ("ach", "growth", "kontribusi"):
                v = _pct(v)
            _put(ws, r, c, v, bold=bold, size=val_size, fmt=fmt, font=_YTD_FONT)

    r = 5
    for m in sheet["months"]:
        write(r, m, False, 10)
        r += 1
    write(r, sheet["total"], True, 10, val_size=10)
    r += 2
    _put(ws, r, 2, "KUARTAL", bold=True, size=10, align="center", border=False, font=_YTD_FONT)
    r += 1
    for q in sheet["quarters"]:
        write(r, q, False, 8, label_align="center")
        r += 1
    write(r, sheet["total"], True, 8, label_align="center", val_size=8)
    if sheet["open"]:
        _put(ws, r + 2, 2, "* dihitung live dari ERP (blm dibekukan): " + ", ".join(f"{m:02d}/{y}" for y, m in sheet["open"]),
             size=8, border=False, font=_YTD_FONT)
    # spt file: ACH < 90% & GROWTH minus merah
    red = Font(color="FFFF0000")
    for _y, key, _fmt, c in layout:
        cl = get_column_letter(c)
        cond = {"ach": "<0.9", "growth": "<0"}.get(key)
        if cond:
            ws.conditional_formatting.add(f"{cl}5:{cl}{r}", FormulaRule(formula=[f"AND(ISNUMBER({cl}5),{cl}5{cond})"], font=red))
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 22.3
    for y, key, _fmt, c in layout:
        ws.column_dimensions[get_column_letter(c)].width = {"target": 14.3, "sales": 14.3, "ach": 7.5, "growth": 8.5,
                                                            "kontribusi": 9}[key]
    ws.freeze_panes = "C5"


def build_sales_ytd(sheets):
    """sheets = [(kode sheet, judul, hasil sales_ytd.build_sheet)] -> 1 sheet Excel per kode (spt file)."""
    wb = Workbook()
    wb.remove(wb.active)
    for kode, title, sheet in sheets:
        _ytd_sheet(wb.create_sheet(kode[:31]), title, sheet)
    return wb


@frappe.whitelist()
def download_sales_ytd(filters=None):
    frappe.only_for("System Manager")
    from frappe.utils import today
    from custom_reports.sales_ytd import YTD, build_sheet, sheet_meta

    filters = _filters(filters)
    upto = int(filters.get("tahun") or getdate(today()).year)
    order, titles = sheet_meta()
    ytd = YTD()
    sheets = [(k, titles.get(k, k), build_sheet(ytd, k, upto)) for k in order]
    _send(build_sales_ytd(sheets), f"SALES X-SHA GROUP_YTD_ {upto}.xlsx")
